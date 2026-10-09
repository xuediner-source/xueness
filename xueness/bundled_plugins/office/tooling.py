"""Bounded structured authoring tools for common Office document formats."""
from __future__ import annotations

import hashlib
import hmac
import io
import math
import os
import re
import stat
import tempfile
import zipfile
import zlib
from pathlib import Path, PurePosixPath, PureWindowsPath

from ...resources import _is_link, replace_file
from ...tool_contract import BuiltinTool
from ...write_lock import DEFAULT_LOCKS, owner_for

MAX_FILE_BYTES = 8_000_000
MAX_ARCHIVE_ENTRIES = 1_200
MAX_ARCHIVE_UNCOMPRESSED = 32_000_000
MAX_XML_PART_BYTES = 4_000_000
MAX_DOCUMENT_CHARS = 400_000
MAX_TEXT_CHARS = 8_000
MAX_READ_PROOFS = 20
MAX_PDF_PAGES = 100
MAX_PDF_PAGE_STREAM_BYTES = 2_000_000
MAX_PDF_TOTAL_STREAM_BYTES = 16_000_000
MAX_SLIDES = 100
MAX_SHEETS = 20
MAX_ROWS = 500
MAX_COLUMNS = 200
MAX_XLSX_CELLS = 100_000
FORMATS = {"docx", "pptx", "xlsx", "pdf"}
_SHA256 = re.compile(r"[0-9a-f]{64}\Z")
_WORD_STYLES = {"Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3",
                "Heading 4", "Heading 5", "Heading 6", "List Bullet", "List Number"}


def _gate(gate, kind: str, path: str, call_id: str | None) -> None:
    # The exact path is the Gate subject and binds web approvals to this call.
    if getattr(gate, "web_approval_gate", False) and kind in ("write", "edit"):
        gate.check(kind, path, call_id)
    else:
        gate.check(kind, path)


def _target(root: Path, path: str, *, must_exist: bool) -> Path:
    """Resolve a workspace-relative path without following a symlink component."""
    if (not isinstance(path, str) or not path or len(path) > 2_000
            or "\0" in path or "\\" in path):
        raise ValueError("path must be a workspace-relative path")
    windows = PureWindowsPath(path)
    if (Path(path).is_absolute() or windows.is_absolute() or windows.drive
            or path.startswith("/") or any(part in ("", ".", "..") for part in path.split("/"))):
        raise ValueError("path must be a workspace-relative path without traversal")
    workspace = Path(root).resolve()
    if not workspace.is_dir():
        raise ValueError("workspace is unavailable")
    parts = path.split("/")
    current = workspace
    for part in parts[:-1]:
        current = current / part
        if _is_link(current):
            raise PermissionError("office paths must not cross symlinks or reparse points")
        try:
            info = current.lstat()
        except FileNotFoundError:
            raise ValueError("parent directory must already exist") from None
        if not stat.S_ISDIR(info.st_mode):
            raise ValueError("parent path is not a directory")
    target = current / parts[-1]
    if _is_link(target):
        raise PermissionError("office files must not be symlinks or reparse points")
    try:
        info = target.lstat()
    except FileNotFoundError:
        info = None
    if info is not None and not stat.S_ISREG(info.st_mode):
        raise ValueError("office path must be a regular file")
    if must_exist and info is None:
        raise ValueError("office file not found")
    resolved = target.resolve(strict=False)
    if not resolved.is_relative_to(workspace):
        raise PermissionError("path outside workspace")
    return target


def _read_bytes(target: Path) -> bytes:
    flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0)
    fd = os.open(target, flags)
    try:
        info = os.fstat(fd)
        if not stat.S_ISREG(info.st_mode):
            raise ValueError("office path must be a regular file")
        if info.st_size > MAX_FILE_BYTES:
            raise ValueError("office file exceeds the 8 MB read limit")
        with os.fdopen(fd, "rb", closefd=False) as stream:
            data = stream.read(MAX_FILE_BYTES + 1)
        if len(data) > MAX_FILE_BYTES:
            raise ValueError("office file exceeds the 8 MB read limit")
        return data
    finally:
        os.close(fd)


def _safe_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("Office archive exceeds the entry limit")
            seen = set()
            total = 0
            actual_total = 0
            for info in infos:
                name = info.filename.replace("\\", "/")
                path = PurePosixPath(name)
                parts = name.rstrip("/").split("/")
                if (not name or name.startswith("/") or path.is_absolute()
                        or any(part in ("", ".", "..") for part in parts)
                        or name in seen or info.flag_bits & 1):
                    raise ValueError("Office archive contains an unsafe member")
                seen.add(name)
                total += info.file_size
                if total > MAX_ARCHIVE_UNCOMPRESSED:
                    raise ValueError("Office archive exceeds the decompressed size limit")
                if name.lower().endswith((".xml", ".rels")) and info.file_size > MAX_XML_PART_BYTES:
                    raise ValueError("Office XML part exceeds the read limit")
                if info.is_dir():
                    continue
                actual_size = 0
                with archive.open(info) as stream:
                    while chunk := stream.read(65_536):
                        actual_size += len(chunk)
                        actual_total += len(chunk)
                        if (actual_size > MAX_ARCHIVE_UNCOMPRESSED
                                or actual_total > MAX_ARCHIVE_UNCOMPRESSED):
                            raise ValueError("Office archive exceeds the decompressed size limit")
                        if (name.lower().endswith((".xml", ".rels"))
                                and actual_size > MAX_XML_PART_BYTES):
                            raise ValueError("Office XML part exceeds the read limit")
                if actual_size != info.file_size:
                    raise ValueError("Office archive member size does not match its directory")
    except (zipfile.BadZipFile, OSError, RuntimeError, EOFError, NotImplementedError) as exc:
        raise ValueError("invalid Office archive") from exc
    except zlib.error as exc:
        raise ValueError("invalid Office archive") from exc


def _record_read(session, path: str, fmt: str, digest: str) -> None:
    if not isinstance(session, dict):
        return
    proofs = session.get("_office_read_proofs")
    if not isinstance(proofs, dict):
        proofs = {}
        session["_office_read_proofs"] = proofs
    proofs[path] = {"format": fmt, "sha256": digest}
    while len(proofs) > MAX_READ_PROOFS:
        proofs.pop(next(iter(proofs)))


def _clear_read(session, path: str) -> None:
    proofs = session.get("_office_read_proofs") if isinstance(session, dict) else None
    if isinstance(proofs, dict):
        proofs.pop(path, None)


def _read_office(data: bytes, fmt: str) -> tuple[dict, bool]:
    from .office_preview import preview_bytes
    _safe_zip(data)
    preview = preview_bytes(data, "." + fmt, include_full_document=False)
    sections = preview.get("sections", [])
    truncation = [bool(preview.get("truncated") or preview.get("imageCount"))]

    def clip(value, limit=MAX_TEXT_CHARS):
        text = str(value)
        if len(text) > limit:
            truncation[0] = True
        return text[:limit]

    if fmt == "docx":
        paragraphs, tables = [], []
        if len(sections) > 1:
            truncation[0] = True
        for section in sections[:1]:
            for block in section.get("blocks", []):
                if block.get("type") == "table":
                    raw_rows = block.get("table", [])
                    if len(raw_rows) > 200 or any(len(row) > 50 for row in raw_rows[:200]):
                        truncation[0] = True
                    tables.append({"rows": [[clip(cell) for cell in row[:50]]
                                              for row in raw_rows[:200]]})
                elif block.get("type") == "paragraph":
                    row = {"text": clip(block.get("text", ""))}
                    style = block.get("style", {}).get("paragraphStyle")
                    if isinstance(style, str) and style in _WORD_STYLES:
                        row["style"] = style
                    paragraphs.append(row)
        return {"paragraphs": paragraphs, "tables": tables}, truncation[0]
    if fmt == "pptx":
        slides = []
        if len(sections) > MAX_SLIDES:
            truncation[0] = True
        for section in sections[:MAX_SLIDES]:
            blocks = [block for block in section.get("blocks", []) if block.get("text")]
            texts = [clip(block.get("text", "")) for block in blocks]
            title_index = next((index for index, block in enumerate(blocks)
                                if int(block.get("position", {}).get("y", 2_000_000)) < 1_000_000), None)
            slide = {"paragraphs": [text for index, text in enumerate(texts) if index != title_index]}
            if title_index is not None:
                slide["title"] = texts[title_index]
            slides.append(slide)
        return {"slides": slides}, truncation[0]
    if fmt == "xlsx":
        sheets = []
        if len(sections) > MAX_SHEETS:
            truncation[0] = True
        for section in sections[:MAX_SHEETS]:
            block = next((item for item in section.get("blocks", []) if item.get("type") == "table"), None)
            if block is None:
                continue
            raw_rows = block.get("table", [])
            if len(raw_rows) > MAX_ROWS or any(len(row) > MAX_COLUMNS for row in raw_rows[:MAX_ROWS]):
                truncation[0] = True
            rows = [[clip(value) for value in row[:MAX_COLUMNS]]
                    for row in raw_rows[:MAX_ROWS]]
            sheets.append({"name": clip(section.get("name", "Sheet"), 31), "rows": rows})
        return {"sheets": sheets}, truncation[0]
    raise ValueError("unsupported Office format")


def _read_pdf(data: bytes) -> tuple[dict, bool]:
    try:
        from pypdf import PdfReader, apply_configuration
    except (ImportError, OSError) as exc:
        raise ValueError("PDF reading requires the Office dependencies in desktop/requirements-office.txt") from exc
    try:
        with apply_configuration(
                maximum_declared_stream_length=MAX_PDF_PAGE_STREAM_BYTES,
                array_based_stream_maximum_output_length=MAX_PDF_PAGE_STREAM_BYTES,
                jbig2_maximum_output_length=MAX_PDF_PAGE_STREAM_BYTES,
                lzw_maximum_output_length=MAX_PDF_PAGE_STREAM_BYTES,
                run_length_maximum_output_length=MAX_PDF_PAGE_STREAM_BYTES,
                zlib_maximum_output_length=MAX_PDF_PAGE_STREAM_BYTES,
                zlib_maximum_recovery_input_length=1_000_000,
                flate_maximum_columns=20_000,
                flate_maximum_row_length=500_000,
                image_maximum_buffer_size=MAX_PDF_PAGE_STREAM_BYTES,
                xmp_maximum_input_length=1_000_000,
                xmp_maximum_element_count=20_000,
                outline_maximum_entries=10_000,
                outline_maximum_depth=40,
                page_tree_maximum_entries=1_000,
                page_tree_maximum_depth=40,
                xform_maximum_invocations_per_extraction=100):
            reader = PdfReader(io.BytesIO(data), strict=True, root_object_recovery_limit=10_000)
            if reader.is_encrypted:
                raise ValueError("encrypted PDFs are not supported")
            count = len(reader.pages)
            pages = []
            remaining = MAX_DOCUMENT_CHARS
            decoded_total = 0
            text_truncated = False
            for index, page in enumerate(reader.pages[:MAX_PDF_PAGES]):
                contents = page.get_contents()
                if contents is not None:
                    decoded_size = len(contents.get_data())
                    decoded_total += decoded_size
                    if (decoded_size > MAX_PDF_PAGE_STREAM_BYTES
                            or decoded_total > MAX_PDF_TOTAL_STREAM_BYTES):
                        raise ValueError("PDF page content exceeds the decompressed read limit")
                text = page.extract_text() or ""
                clipped = text[:min(10_000, remaining)]
                text_truncated |= len(clipped) < len(text)
                pages.append({"number": index + 1, "text": clipped})
                remaining -= len(clipped)
                if remaining <= 0:
                    break
            metadata = reader.metadata
            raw_title = str(metadata.title) if metadata and metadata.title else ""
            title = raw_title[:200]
            text_truncated |= len(title) < len(raw_title)
            return {"title": title, "pages": pages}, (
                count > len(pages) or remaining <= 0 or text_truncated
            )
    except ValueError:
        raise
    except Exception as exc:  # malformed parser input is a bounded tool failure
        raise ValueError("invalid PDF document") from exc


def delivery_content_text(data: bytes, suffix: str) -> str:
    """Return only bounded, parsed visible text for a delivery check.

    The completion checker needs the same safe parsers as ``office_read``, but
    must not match schema labels, worksheet names, or PDF metadata as document
    content. Any parser truncation makes the check inconclusive and therefore
    fails closed.
    """
    if not isinstance(data, bytes) or len(data) > MAX_FILE_BYTES:
        raise ValueError("Office file exceeds the bounded read limit")
    if not isinstance(suffix, str):
        raise ValueError("unsupported Office format")
    fmt = suffix.lower().lstrip(".")
    if fmt not in FORMATS:
        raise ValueError("unsupported Office format")

    document, truncated = _read_pdf(data) if fmt == "pdf" else _read_office(data, fmt)
    if truncated:
        raise ValueError("Office content is truncated and cannot be checked")

    visible = []
    visible_chars = [0]

    def add(value):
        if not isinstance(value, str):
            raise ValueError("invalid parsed Office text")
        if value:
            next_size = visible_chars[0] + len(value) + bool(visible)
            if next_size > MAX_DOCUMENT_CHARS:
                raise ValueError("Office content exceeds the bounded text limit")
            visible.append(value)
            visible_chars[0] = next_size

    if fmt == "docx":
        for paragraph in document.get("paragraphs", []):
            add(paragraph.get("text"))
        for table in document.get("tables", []):
            for row in table.get("rows", []):
                for cell in row:
                    add(cell)
    elif fmt == "pptx":
        for slide in document.get("slides", []):
            title = slide.get("title")
            if title is not None:
                add(title)
            for paragraph in slide.get("paragraphs", []):
                add(paragraph)
    elif fmt == "xlsx":
        for sheet in document.get("sheets", []):
            for row in sheet.get("rows", []):
                for cell in row:
                    add(cell)
    else:  # PDF title is metadata; only extracted page text is document content.
        for page in document.get("pages", []):
            add(page.get("text"))
    return "\n".join(visible)


def _read(root, gate, args, session, call_id):
    path = args.get("path")
    if not isinstance(path, str):
        raise ValueError("path must be a string")
    _gate(gate, "read", path, call_id)
    target = _target(root, path, must_exist=True)
    fmt = target.suffix.lower().lstrip(".")
    if fmt not in FORMATS:
        raise ValueError("supported formats are DOCX, PPTX, XLSX and PDF")
    data = _read_bytes(target)
    digest = hashlib.sha256(data).hexdigest()
    if fmt == "pdf":
        document, truncated = _read_pdf(data)
    else:
        document, truncated = _read_office(data, fmt)
    if truncated:
        _clear_read(session, path)
    else:
        _record_read(session, path, fmt, digest)
    return {"ok": True, "path": path, "format": fmt, "sha256": digest,
            "document": document, "truncated": truncated}


def _budget_text(value: str, field: str, budget: list[int], *, allow_empty=True) -> str:
    if not isinstance(value, str) or "\0" in value or len(value) > MAX_TEXT_CHARS:
        raise ValueError(f"{field} must be text of at most {MAX_TEXT_CHARS} characters")
    if not allow_empty and not value.strip():
        raise ValueError(f"{field} must not be empty")
    budget[0] += len(value)
    if budget[0] > MAX_DOCUMENT_CHARS:
        raise ValueError("document text exceeds the 400000 character limit")
    return value


def _only_keys(item: dict, allowed: set[str], field: str) -> None:
    if not isinstance(item, dict) or set(item) - allowed:
        raise ValueError(f"invalid {field} fields")


def _document_object(document, fmt: str) -> dict:
    if not isinstance(document, dict):
        raise ValueError("document must be an object")
    budget = [0]
    title = _budget_text(document.get("title", ""), "title", budget)
    if fmt == "docx":
        _only_keys(document, {"title", "paragraphs", "tables"}, "DOCX document")
        paragraphs = document.get("paragraphs", [])
        tables = document.get("tables", [])
        if not isinstance(paragraphs, list) or len(paragraphs) > 500:
            raise ValueError("DOCX paragraphs must be a list of at most 500 items")
        if not isinstance(tables, list) or len(tables) > 100:
            raise ValueError("DOCX tables must be a list of at most 100 items")
        checked_paragraphs = []
        for item in paragraphs:
            _only_keys(item, {"text", "style", "bold", "italic"}, "paragraph")
            text = _budget_text(item.get("text"), "paragraph text", budget)
            style = item.get("style", "Normal")
            if not isinstance(style, str) or style not in _WORD_STYLES:
                raise ValueError("unsupported DOCX paragraph style")
            if type(item.get("bold", False)) is not bool or type(item.get("italic", False)) is not bool:
                raise ValueError("paragraph bold and italic values must be boolean")
            checked_paragraphs.append({"text": text, "style": style,
                                      "bold": item.get("bold", False),
                                      "italic": item.get("italic", False)})
        checked_tables = []
        for table in tables:
            _only_keys(table, {"rows"}, "table")
            rows = table.get("rows")
            if not isinstance(rows, list) or not rows or len(rows) > 200:
                raise ValueError("a DOCX table needs 1 to 200 rows")
            checked_rows = []
            width = 0
            for row in rows:
                if not isinstance(row, list) or len(row) > 50:
                    raise ValueError("DOCX table rows must have at most 50 cells")
                width = max(width, len(row))
                checked_rows.append([_budget_text(value, "table cell", budget) for value in row])
            if not width:
                raise ValueError("DOCX table rows must contain cells")
            checked_tables.append({"rows": checked_rows})
        if not title and not checked_paragraphs and not checked_tables:
            raise ValueError("DOCX document needs a title, paragraph or table")
        return {"title": title, "paragraphs": checked_paragraphs, "tables": checked_tables}
    if fmt == "pptx":
        _only_keys(document, {"title", "slides"}, "PPTX document")
        slides = document.get("slides")
        if not isinstance(slides, list) or not 1 <= len(slides) <= MAX_SLIDES:
            raise ValueError("PPTX document needs 1 to 100 slides")
        checked_slides = []
        for slide in slides:
            _only_keys(slide, {"title", "paragraphs", "bullets"}, "slide")
            slide_title = _budget_text(slide.get("title", ""), "slide title", budget)
            if len(slide_title) > 60:
                raise ValueError("slide titles are limited to 60 characters")
            paragraphs = slide.get("paragraphs", [])
            bullets = slide.get("bullets", [])
            if (not isinstance(paragraphs, list) or len(paragraphs) > 100
                    or not isinstance(bullets, list) or len(bullets) > 100):
                raise ValueError("slide paragraphs and bullets must each have at most 100 items")
            if len(paragraphs) + len(bullets) > 20:
                raise ValueError("a slide can have at most 20 paragraphs and bullets")
            checked_paragraphs = [_budget_text(value, "slide paragraph", budget) for value in paragraphs]
            checked_bullets = [_budget_text(value, "slide bullet", budget) for value in bullets]
            if any(len(value) > 400 for value in checked_paragraphs + checked_bullets):
                raise ValueError("slide paragraphs and bullets are limited to 400 characters each")
            display_lines = sum(max(1, math.ceil(len(value) / 50))
                                for value in checked_paragraphs + checked_bullets)
            if display_lines > 18:
                raise ValueError("slide content exceeds the bounded text layout")
            checked_slides.append({
                "title": slide_title,
                "paragraphs": checked_paragraphs,
                "bullets": checked_bullets,
            })
        return {"title": title, "slides": checked_slides}
    if fmt == "xlsx":
        _only_keys(document, {"title", "sheets"}, "XLSX document")
        sheets = document.get("sheets")
        if not isinstance(sheets, list) or not 1 <= len(sheets) <= MAX_SHEETS:
            raise ValueError("XLSX document needs 1 to 20 sheets")
        checked_sheets, names = [], set()
        total_cells = 0
        for sheet in sheets:
            _only_keys(sheet, {"name", "rows"}, "sheet")
            name = sheet.get("name")
            if (not isinstance(name, str) or not name.strip() or len(name) > 31
                    or any(char in name for char in "[]:*?/\\") or name.startswith("'") or name.endswith("'")):
                raise ValueError("invalid worksheet name")
            if name.casefold() in names:
                raise ValueError("worksheet names must be unique")
            names.add(name.casefold())
            _budget_text(name, "worksheet name", budget)
            rows = sheet.get("rows")
            if not isinstance(rows, list) or len(rows) > MAX_ROWS:
                raise ValueError("worksheet rows must be a list of at most 500 rows")
            checked_rows = []
            for row in rows:
                if not isinstance(row, list) or len(row) > MAX_COLUMNS:
                    raise ValueError("worksheet rows must have at most 200 cells")
                total_cells += len(row)
                if total_cells > MAX_XLSX_CELLS:
                    raise ValueError("XLSX document exceeds 100000 populated cell slots")
                checked = []
                for value in row:
                    if value is None or type(value) is bool:
                        checked.append(value)
                    elif type(value) is int:
                        if abs(value) > 999_999_999_999_999:
                            raise ValueError("XLSX integers are limited to 15 digits")
                        checked.append(value)
                    elif type(value) is float:
                        if not math.isfinite(value) or abs(value) > 1e100:
                            raise ValueError("XLSX numbers must be finite and within range")
                        checked.append(value)
                    elif isinstance(value, str):
                        checked.append(_budget_text(value, "worksheet cell", budget))
                    else:
                        raise ValueError("XLSX cells must be text, number, boolean or null")
                checked_rows.append(checked)
            checked_sheets.append({"name": name, "rows": checked_rows})
        return {"title": title, "sheets": checked_sheets}
    if fmt == "pdf":
        _only_keys(document, {"title", "paragraphs"}, "PDF document")
        paragraphs = document.get("paragraphs", [])
        if not isinstance(paragraphs, list) or len(paragraphs) > 500:
            raise ValueError("PDF paragraphs must be a list of at most 500 items")
        checked = []
        for item in paragraphs:
            _only_keys(item, {"text", "style"}, "PDF paragraph")
            text = _budget_text(item.get("text"), "PDF paragraph", budget)
            style = item.get("style", "body")
            if style not in ("title", "heading", "body"):
                raise ValueError("PDF style must be title, heading or body")
            checked.append({"text": text, "style": style})
        if not title and not checked:
            raise ValueError("PDF document needs a title or paragraph")
        return {"title": title, "paragraphs": checked}
    raise ValueError("unsupported Office format")


def _make_docx(document) -> bytes:
    from docx import Document
    result = Document()
    if document["title"]:
        result.add_heading(document["title"], level=0)
    for item in document["paragraphs"]:
        paragraph = result.add_paragraph(style=item["style"])
        run = paragraph.add_run(item["text"])
        run.bold = item["bold"]
        run.italic = item["italic"]
    for item in document["tables"]:
        width = max(len(row) for row in item["rows"])
        table = result.add_table(rows=0, cols=width)
        table.style = "Table Grid"
        for values in item["rows"]:
            row = table.add_row().cells
            for index, value in enumerate(values):
                row[index].text = value
    result.core_properties.title = document["title"]
    stream = io.BytesIO()
    result.save(stream)
    return stream.getvalue()


def _make_pptx(document) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    result = Presentation()
    for item in document["slides"]:
        slide = result.slides.add_slide(result.slide_layouts[6])
        title = item["title"]
        if title:
            box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35),
                                           result.slide_width - Inches(1.1), Inches(0.9))
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = title
            paragraph.font.size = Pt(26)
            paragraph.font.bold = True
        blocks = [(text, False) for text in item["paragraphs"]]
        blocks.extend((text, True) for text in item["bullets"])
        y = 1.5
        for text, bullet in blocks:
            lines = max(1, math.ceil(len(text) / 50))
            height = min(1.5, 0.28 * lines)
            box = slide.shapes.add_textbox(Inches(0.7), Inches(y),
                                           result.slide_width - Inches(1.4), Inches(height))
            box.text_frame.word_wrap = True
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = ("• " if bullet else "") + text
            paragraph.font.size = Pt(16)
            y += height + 0.04
            if y > 7.15:
                raise ValueError("slide content exceeds the bounded text layout")
    result.core_properties.title = document["title"]
    stream = io.BytesIO()
    result.save(stream)
    return stream.getvalue()


def _make_xlsx(document) -> bytes:
    from openpyxl import Workbook
    result = Workbook()
    result.remove(result.active)
    result.properties.title = document["title"]
    for sheet_data in document["sheets"]:
        sheet = result.create_sheet(sheet_data["name"])
        for row_index, row in enumerate(sheet_data["rows"], 1):
            for column_index, value in enumerate(row, 1):
                cell = sheet.cell(row=row_index, column=column_index, value=value)
                # openpyxl infers formulas from strings starting with '='.
                # Force every model-provided string to a literal text cell.
                if isinstance(value, str):
                    cell.data_type = "s"
    stream = io.BytesIO()
    result.save(stream)
    return stream.getvalue()


def _make_pdf(document) -> bytes:
    from xml.sax.saxutils import escape
    from reportlab.lib import pagesizes
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except KeyError:
        pass
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("OfficeTitle", parent=styles["Title"], fontName="STSong-Light")
    heading_style = ParagraphStyle("OfficeHeading", parent=styles["Heading2"], fontName="STSong-Light")
    body_style = ParagraphStyle("OfficeBody", parent=styles["BodyText"], fontName="STSong-Light",
                                leading=16)
    stream = io.BytesIO()
    result = SimpleDocTemplate(stream, pagesize=pagesizes.A4, title=document["title"] or "Office document",
                               pageCompression=1)
    story = []
    if document["title"]:
        story.extend([Paragraph(escape(document["title"]), title_style), Spacer(1, 12)])
    for item in document["paragraphs"]:
        style = {"title": title_style, "heading": heading_style, "body": body_style}[item["style"]]
        text = escape(item["text"]).replace("\n", "<br/>")
        story.extend([Paragraph(text or " ", style), Spacer(1, 6)])
    result.build(story)
    return stream.getvalue()


def _build_document(fmt: str, document) -> bytes:
    checked = _document_object(document, fmt)
    try:
        data = {"docx": _make_docx, "pptx": _make_pptx,
                "xlsx": _make_xlsx, "pdf": _make_pdf}[fmt](checked)
    except (ImportError, OSError) as exc:
        raise ValueError("Office authoring requires the dependencies in desktop/requirements-office.txt") from exc
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("generated document exceeds the 8 MB file limit")
    return data


def _write_atomic(target: Path, data: bytes, *, create_only: bool, root: Path, path: str,
                  expected_sha256: str | None = None) -> None:
    """Write a complete artifact next to its target and atomically install it."""
    parent = target.parent
    if _is_link(parent) or not parent.is_dir():
        raise PermissionError("office parent directory must be a real workspace directory")
    fd, temporary = tempfile.mkstemp(prefix=".xueness-office-", suffix=".tmp", dir=parent)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        _target(root, path, must_exist=not create_only)
        if create_only:
            # Hard-link installation is atomic and fails instead of replacing a
            # file that appeared after the create-only check.
            os.link(temporary, target, follow_symlinks=False)
            os.unlink(temporary)
        else:
            # Windows sharing collisions retry inside replace_file. Recheck the
            # original read proof before every attempt, including the first, so
            # an intervening edit is never silently replaced.
            def _prove_unchanged():
                _target(root, path, must_exist=True)
                if expected_sha256 and not hmac.compare_digest(
                        hashlib.sha256(_read_bytes(target)).hexdigest(), expected_sha256):
                    raise ValueError("office file changed after office_read; read it again before replacing")

            replace_file(temporary, target, before_replace=_prove_unchanged)
        try:
            directory_fd = os.open(parent, os.O_RDONLY | getattr(os, "O_DIRECTORY", 0))
        except OSError:
            directory_fd = None
        if directory_fd is not None:
            try:
                os.fsync(directory_fd)
            finally:
                os.close(directory_fd)
    finally:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass


def _format_for(path: str, fmt) -> str:
    if not isinstance(fmt, str) or fmt not in FORMATS:
        raise ValueError("format must be docx, pptx, xlsx or pdf")
    if Path(path).suffix.lower() != "." + fmt:
        raise ValueError("format must match the file extension")
    return fmt


def _create(root, gate, args, session, call_id):
    path = args.get("path")
    fmt = _format_for(path, args.get("format")) if isinstance(path, str) else None
    if fmt is None:
        raise ValueError("path must be a string")
    _gate(gate, "write", path, call_id)
    target = _target(root, path, must_exist=False)
    if target.exists():
        raise ValueError("office_create only creates a new file; use office_read then office_replace")
    data = _build_document(fmt, args.get("document"))
    owner = owner_for(session)
    if not DEFAULT_LOCKS.acquire(target, owner):
        return {"ok": False, "conflict": True, "error": "path busy"}
    try:
        if target.exists() or _is_link(target):
            raise ValueError("target appeared before document creation")
        _write_atomic(target, data, create_only=True, root=root, path=path)
    finally:
        DEFAULT_LOCKS.release(target, owner)
    digest = hashlib.sha256(data).hexdigest()
    return {"ok": True, "path": path, "format": fmt, "bytes": len(data), "sha256": digest}


def _replace(root, gate, args, session, call_id):
    path = args.get("path")
    fmt = _format_for(path, args.get("format")) if isinstance(path, str) else None
    if fmt is None:
        raise ValueError("path must be a string")
    expected = args.get("expected_sha256")
    if not isinstance(expected, str) or not _SHA256.fullmatch(expected):
        raise ValueError("expected_sha256 from office_read is required")
    _gate(gate, "edit", path, call_id)
    target = _target(root, path, must_exist=True)
    proofs = session.get("_office_read_proofs") if isinstance(session, dict) else None
    proof = proofs.get(path) if isinstance(proofs, dict) else None
    if (not isinstance(proof, dict) or proof.get("format") != fmt
            or not isinstance(proof.get("sha256"), str)
            or not hmac.compare_digest(expected, proof["sha256"])):
        raise PermissionError("read this exact path in the current session before replacing it")
    existing = _read_bytes(target)
    current = hashlib.sha256(existing).hexdigest()
    if not hmac.compare_digest(current, expected):
        raise ValueError("office file changed after office_read; read it again before replacing")
    data = _build_document(fmt, args.get("document"))
    owner = owner_for(session)
    if not DEFAULT_LOCKS.acquire(target, owner):
        return {"ok": False, "conflict": True, "error": "path busy"}
    try:
        # Recheck under the writer lock immediately before atomic replacement.
        _target(root, path, must_exist=True)
        latest = hashlib.sha256(_read_bytes(target)).hexdigest()
        if not hmac.compare_digest(latest, expected):
            raise ValueError("office file changed after office_read; read it again before replacing")
        _write_atomic(target, data, create_only=False, root=root, path=path,
                      expected_sha256=expected)
    finally:
        DEFAULT_LOCKS.release(target, owner)
    proofs.pop(path, None)
    digest = hashlib.sha256(data).hexdigest()
    return {"ok": True, "path": path, "format": fmt, "bytes": len(data), "sha256": digest}


_PARAGRAPH_SCHEMA = {"type": "array", "items": {"type": "object",
    "properties": {"text": {"type": "string"}, "style": {"type": "string"},
                   "bold": {"type": "boolean"}, "italic": {"type": "boolean"}},
    "required": ["text"], "additionalProperties": False}}
_TABLE_SCHEMA = {"type": "array", "items": {"type": "object", "properties": {
    "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}}},
    "required": ["rows"], "additionalProperties": False}}
_SLIDE_SCHEMA = {"type": "array", "items": {"type": "object", "properties": {
    "title": {"type": "string"}, "paragraphs": {"type": "array", "items": {"type": "string"}},
    "bullets": {"type": "array", "items": {"type": "string"}}}, "additionalProperties": False}}
_SHEET_SCHEMA = {"type": "array", "items": {"type": "object", "properties": {
    "name": {"type": "string"}, "rows": {"type": "array", "items": {
        "type": "array", "items": {"type": ["string", "number", "boolean", "null"]}}}},
    "required": ["name", "rows"], "additionalProperties": False}}
_DOCUMENT_SCHEMA = {"type": "object", "properties": {
    "title": {"type": "string"}, "paragraphs": _PARAGRAPH_SCHEMA,
    "tables": _TABLE_SCHEMA, "slides": _SLIDE_SCHEMA, "sheets": _SHEET_SCHEMA},
    "additionalProperties": False}

REGISTRY = (
    BuiltinTool("office_read", "Read bounded structured text from a DOCX, PPTX, XLSX or PDF in the workspace",
                {"path": {"type": "string", "description": "Workspace-relative Office file path"}},
                ("path",), "read", False, _read, concurrency_safe=True),
    BuiltinTool("office_create", "Create a structured DOCX, PPTX, XLSX or PDF in the workspace (approval required)",
                {"path": {"type": "string"}, "format": {"type": "string", "enum": ["docx", "pptx", "xlsx", "pdf"]},
                 "document": _DOCUMENT_SCHEMA},
                ("path", "format", "document"), "write", True, _create,
                approval_subject=lambda args: args.get("path", "")),
    BuiltinTool("office_replace", "Replace a structured Office document after reading the same path and supplying its sha256 (approval required)",
                {"path": {"type": "string"}, "format": {"type": "string", "enum": ["docx", "pptx", "xlsx", "pdf"]},
                 "expected_sha256": {"type": "string", "description": "sha256 returned by office_read for this exact path in the current session"},
                 "document": _DOCUMENT_SCHEMA},
                ("path", "format", "expected_sha256", "document"), "edit", True, _replace,
                approval_subject=lambda args: args.get("path", "")),
)
