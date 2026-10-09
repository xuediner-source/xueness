"""Focused behavior and safety checks for Office authoring tools."""
from __future__ import annotations

import hashlib
import io
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from xueness.bundled_plugins.office import tooling


class Gate:
    web_approval_gate = True

    def __init__(self):
        self.calls = []

    def check(self, kind, subject, call_id=None):
        self.calls.append((kind, subject, call_id))


class OfficeToolTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.session = {"id": "office-test-session"}
        self.gate = Gate()

    def tearDown(self):
        self.temp.cleanup()

    def create(self, path, fmt, document):
        result = tooling._create(self.root, self.gate,
                                 {"path": path, "format": fmt, "document": document},
                                 self.session, "write-call")
        self.assertTrue(result["ok"])
        return result

    def read(self, path):
        result = tooling._read(self.root, self.gate, {"path": path},
                               self.session, "read-call")
        self.assertTrue(result["ok"])
        return result

    def test_create_and_read_all_supported_formats(self):
        docx = self.create("brief.docx", "docx", {
            "title": "Project Brief",
            "paragraphs": [{"text": "Summary", "style": "Heading 1"},
                           {"text": "A short report", "bold": True}],
            "tables": [{"rows": [["Item", "Value"], ["Status", "Ready"]]}],
        })
        self.assertEqual(self.read("brief.docx")["document"]["paragraphs"][0]["text"], "Project Brief")
        self.assertEqual(self.read("brief.docx")["document"]["tables"][0]["rows"][1], ["Status", "Ready"])
        self.assertEqual(docx["sha256"], hashlib.sha256((self.root / "brief.docx").read_bytes()).hexdigest())

        self.create("slides.pptx", "pptx", {"title": "Quarterly", "slides": [
            {"title": "Results", "paragraphs": ["Revenue grew"], "bullets": ["North: 12%"]},
        ]})
        pptx = self.read("slides.pptx")["document"]["slides"][0]
        self.assertEqual(pptx["title"], "Results")
        self.assertIn("Revenue grew", pptx["paragraphs"])
        self.assertIn("• North: 12%", pptx["paragraphs"])

        self.create("data.xlsx", "xlsx", {"title": "Metrics", "sheets": [
            {"name": "Data", "rows": [["Label", "Value"], ["Total", 42], ["Literal", "=1+1"]]},
        ]})
        xlsx = self.read("data.xlsx")["document"]["sheets"][0]
        self.assertEqual(xlsx["name"], "Data")
        self.assertEqual(xlsx["rows"][1], ["Total", "42"])
        self.assertEqual(xlsx["rows"][2], ["Literal", "=1+1"])
        with zipfile.ZipFile(self.root / "data.xlsx") as archive:
            sheet_xml = archive.read("xl/worksheets/sheet1.xml")
        self.assertNotIn(b"<f>", sheet_xml)

        self.create("note.pdf", "pdf", {"title": "Meeting Note", "paragraphs": [
            {"text": "Decisions", "style": "heading"},
            {"text": "审核完成。 The review is complete."},
        ]})
        pdf = self.read("note.pdf")["document"]
        self.assertEqual(pdf["title"], "Meeting Note")
        self.assertIn("The review is complete.", pdf["pages"][0]["text"])
        self.assertIn("审核完成。", pdf["pages"][0]["text"])

    def test_replace_requires_session_read_and_rejects_stale_digest(self):
        self.create("draft.docx", "docx", {"paragraphs": [{"text": "Old"}]})
        first = self.read("draft.docx")
        args = {"path": "draft.docx", "format": "docx",
                "expected_sha256": first["sha256"],
                "document": {"paragraphs": [{"text": "New"}]}}
        with self.assertRaises(PermissionError):
            tooling._replace(self.root, self.gate, args, {"id": "other"}, "edit-call")
        with self.assertRaises(PermissionError):
            tooling._replace(self.root, self.gate, {**args, "expected_sha256": "0" * 64},
                              self.session, "edit-call")

        (self.root / "draft.docx").write_bytes(
            tooling._build_document("docx", {"paragraphs": [{"text": "External edit"}]})
        )
        with self.assertRaisesRegex(ValueError, "changed after office_read"):
            tooling._replace(self.root, self.gate, args, self.session, "edit-call")

        current = self.read("draft.docx")
        args["expected_sha256"] = current["sha256"]
        result = tooling._replace(self.root, self.gate, args, self.session, "edit-call")
        self.assertTrue(result["ok"])
        self.assertNotIn("draft.docx", self.session["_office_read_proofs"])
        self.assertEqual(self.read("draft.docx")["document"]["paragraphs"][0]["text"], "New")

    @unittest.skipUnless(os.name == 'nt', 'Windows atomic replacement behavior')
    def test_atomic_replace_recovers_from_one_sharing_violation(self):
        target = self.root/'artifact.docx'
        target.write_bytes(b'old')
        error = PermissionError('fixture sharing violation')
        error.winerror = 32
        actual = tooling.os.replace
        with patch.object(tooling.os, 'replace', side_effect=[error, None]) as rename:
            def attempt(source, destination):
                if rename.call_count == 1:
                    raise error
                actual(source, destination)
            rename.side_effect = attempt
            tooling._write_atomic(target, b'new', create_only=False, root=self.root,
                                  path=target.name, expected_sha256=hashlib.sha256(b'old').hexdigest())
            self.assertEqual(2, rename.call_count)
        self.assertEqual(b'new', target.read_bytes())
        self.assertFalse(list(self.root.glob('.xueness-office-*.tmp')))

    @unittest.skipUnless(os.name == 'nt', 'Windows atomic replacement behavior')
    def test_atomic_retry_rejects_an_intervening_external_edit(self):
        target = self.root/'artifact.docx'
        target.write_bytes(b'old')
        error = PermissionError('fixture sharing violation')
        error.winerror = 32
        def occupied(_source, destination):
            Path(destination).write_bytes(b'external edit')
            raise error
        with patch.object(tooling.os, 'replace', side_effect=occupied) as rename:
            with self.assertRaisesRegex(ValueError, 'changed after office_read'):
                tooling._write_atomic(target, b'new', create_only=False, root=self.root,
                                      path=target.name, expected_sha256=hashlib.sha256(b'old').hexdigest())
            self.assertEqual(1, rename.call_count)
        self.assertEqual(b'external edit', target.read_bytes())
        self.assertFalse(list(self.root.glob('.xueness-office-*.tmp')))

    @unittest.skipUnless(os.name == 'nt', 'Windows atomic replacement behavior')
    def test_atomic_retry_is_bounded_and_retains_original_on_denial(self):
        target = self.root/'artifact.docx'
        target.write_bytes(b'old')
        error = PermissionError('fixture access denied')
        error.winerror = 5
        with patch.object(tooling.os, 'replace', side_effect=error) as rename:
            with self.assertRaises(PermissionError):
                tooling._write_atomic(target, b'new', create_only=False, root=self.root,
                                      path=target.name, expected_sha256=hashlib.sha256(b'old').hexdigest())
            self.assertEqual(3, rename.call_count)
        self.assertEqual(b'old', target.read_bytes())
        self.assertFalse(list(self.root.glob('.xueness-office-*.tmp')))

    def test_gate_receives_exact_subject_and_workspace_symlinks_are_rejected(self):
        self.create("safe.docx", "docx", {"paragraphs": [{"text": "ok"}]})
        self.read("safe.docx")
        with self.assertRaises((PermissionError, ValueError)):
            tooling._create(self.root, self.gate,
                            {"path": "../outside.docx", "format": "docx",
                             "document": {"paragraphs": [{"text": "no"}]}},
                            self.session, "unsafe-write")
        self.assertEqual(self.gate.calls[-1], ("write", "../outside.docx", "unsafe-write"))

        outside = self.root.parent / (self.root.name + "-outside.docx")
        outside.write_bytes(b"outside")
        link = self.root / "linked.docx"
        try:
            link.symlink_to(outside)
        except (OSError, NotImplementedError):
            outside.unlink(missing_ok=True)
            self.skipTest("symlinks are unavailable")
        with self.assertRaises(PermissionError):
            tooling._read(self.root, self.gate, {"path": "linked.docx"}, self.session, "read-link")
        self.assertEqual(self.gate.calls[-1], ("read", "linked.docx", None))
        outside.unlink(missing_ok=True)

    def test_zip_member_traversal_is_rejected_before_preview(self):
        raw = io.BytesIO()
        with zipfile.ZipFile(raw, "w") as archive:
            archive.writestr("../outside", "no")
        with self.assertRaisesRegex(ValueError, "unsafe member"):
            tooling._safe_zip(raw.getvalue())

    def test_truncated_read_does_not_authorize_full_document_replacement(self):
        from docx import Document
        document = Document()
        for index in range(501):
            document.add_paragraph(f"Paragraph {index}")
        document.save(self.root / "long.docx")
        result = self.read("long.docx")
        self.assertTrue(result["truncated"])
        self.assertNotIn("long.docx", self.session.get("_office_read_proofs", {}))
        with self.assertRaises(PermissionError):
            tooling._replace(self.root, self.gate, {
                "path": "long.docx", "format": "docx", "expected_sha256": result["sha256"],
                "document": {"paragraphs": [{"text": "replacement"}]},
            }, self.session, "edit-call")

    def test_long_docx_paragraph_is_nonempty_and_reported_truncated(self):
        from docx import Document
        document = Document()
        document.add_paragraph("A" * 12_000)
        document.save(self.root / "long-paragraph.docx")

        result = self.read("long-paragraph.docx")
        paragraph = result["document"]["paragraphs"][0]["text"]
        self.assertEqual(paragraph, "A" * tooling.MAX_TEXT_CHARS)
        self.assertTrue(result["truncated"])
        self.assertNotIn("long-paragraph.docx", self.session.get("_office_read_proofs", {}))

    def test_long_single_pdf_page_is_reported_truncated(self):
        from reportlab.pdfgen.canvas import Canvas
        stream = io.BytesIO()
        pdf = Canvas(stream)
        pdf.drawString(20, 700, "P" * 20_000)
        pdf.showPage()
        pdf.save()
        (self.root / "long-page.pdf").write_bytes(stream.getvalue())

        result = self.read("long-page.pdf")
        page_text = result["document"]["pages"][0]["text"]
        self.assertGreater(len(page_text), 0)
        self.assertEqual(len(page_text), 10_000)
        self.assertTrue(result["truncated"])
        self.assertNotIn("long-page.pdf", self.session.get("_office_read_proofs", {}))

    def test_registry_owns_tools_by_office_and_disabled_switch_blocks_dispatch(self):
        from xueness.plugin_runtime import tool_owner
        from xueness.tool_contract import bind_execution
        from xueness.tool_registry import dispatch
        for name in ("office_read", "office_create", "office_replace"):
            self.assertEqual(tool_owner(name), "office")

        state = self.root / "state"
        state.mkdir()
        (state / "plugin-state.json").write_text(
            '{"apiVersion":1,"enabled":{"office":false}}', encoding="utf-8")
        with bind_execution(state_dir=state):
            result = dispatch(self.root, self.gate, "office_create", {
                "path": "disabled.docx", "format": "docx",
                "document": {"paragraphs": [{"text": "should not run"}]},
            }, self.session)
        self.assertEqual(result, {"ok": False, "error": "plugin disabled"})
        self.assertEqual(self.gate.calls, [])
        self.assertFalse((self.root / "disabled.docx").exists())


if __name__ == "__main__":
    unittest.main()
